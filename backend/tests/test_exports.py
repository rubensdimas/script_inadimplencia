import csv as csv_module
from io import BytesIO, StringIO

from openpyxl import load_workbook

from tests.test_comparisons import (
    _linha_xlsx as _linha_xlsx_comparacao,
)
from tests.test_comparisons import (
    _preparar_par_completo,
    _upload_csv,
    _upload_xlsx,
)
from tests.test_dashboard import _preparar_dashboard
from tests.test_entities import _preparar_entidades


def _linhas_csv(conteudo: bytes) -> list[list[str]]:
    texto = conteudo.decode("utf-8-sig")
    return list(csv_module.reader(StringIO(texto), delimiter=";"))


def _linhas_xlsx(conteudo: bytes) -> list[tuple]:
    pasta = load_workbook(BytesIO(conteudo))
    aba = pasta.active
    return [tuple(linha) for linha in aba.iter_rows(values_only=True)]


# --- dataset=ranking ---------------------------------------------------


def test_exportar_ranking_csv_ordenado_por_valor_total_e_documento_mascarado(cliente):
    antigo, atual, csv_snap = _preparar_dashboard(cliente)

    resposta = cliente.get(
        "/api/exports/ranking",
        params={"format": "csv", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": atual["id"]},
    )
    assert resposta.status_code == 200, resposta.text
    assert resposta.headers["content-type"].startswith("text/csv")
    assert ".csv" in resposta.headers["content-disposition"]
    assert "ranking" in resposta.headers["content-disposition"]

    linhas = _linhas_csv(resposta.content)
    cabecalho, *dados = linhas
    assert "valor_total" in cabecalho
    assert "documento_mascarado" in cabecalho

    idx_nome = cabecalho.index("nome")
    idx_valor = cabecalho.index("valor_total")
    idx_doc = cabecalho.index("documento_mascarado")

    assert [linha[idx_nome] for linha in dados] == [
        "ANA PROFISSIONAL",
        "EMPRESA DEVEDORA LTDA",
        "CARLOS PROFISSIONAL",
    ]
    assert [linha[idx_valor] for linha in dados] == ["400.00", "150.00", "50.00"]
    assert dados[0][idx_doc] == "***.***.***-11"
    assert "11111111111" not in resposta.text


def test_exportar_ranking_xlsx_valores_numericos_e_documento_mascarado(cliente):
    antigo, atual, csv_snap = _preparar_dashboard(cliente)

    resposta = cliente.get(
        "/api/exports/ranking",
        params={"format": "xlsx", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": atual["id"]},
    )
    assert resposta.status_code == 200, resposta.text
    assert resposta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in resposta.headers["content-disposition"]

    linhas = _linhas_xlsx(resposta.content)
    cabecalho = linhas[0]
    idx_valor = cabecalho.index("valor_total")
    idx_doc = cabecalho.index("documento_mascarado")

    primeira_linha_dados = linhas[1]
    assert isinstance(primeira_linha_dados[idx_valor], (int, float))
    assert float(primeira_linha_dados[idx_valor]) == 400.0
    assert primeira_linha_dados[idx_doc] == "***.***.***-11"

    todos_os_valores = "".join(str(v) for linha in linhas for v in linha)
    assert "11111111111" not in todos_os_valores


# --- dataset=comparisons -------------------------------------------------


def test_exportar_comparisons_csv_agrupa_categorias_e_mascara_documentos(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/comparisons",
        params={"format": "csv", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_csv(resposta.content)
    cabecalho, *dados = linhas
    idx_categoria = cabecalho.index("categoria")

    categorias = {linha[idx_categoria] for linha in dados}
    assert categorias == {
        "somente_csv",
        "somente_xlsx",
        "obrigacao_somente_csv",
        "obrigacao_somente_xlsx",
        "conflito_parcelamento",
        "nome_ambiguo",
    }
    assert sum(1 for linha in dados if linha[idx_categoria] == "nome_ambiguo") == 2

    assert "66666666666" not in resposta.text
    assert "77777777777" not in resposta.text


def test_exportar_comparisons_xlsx_representa_conflito_de_parcelamento(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/comparisons",
        params={"format": "xlsx", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_xlsx(resposta.content)
    cabecalho, *dados = linhas
    idx_categoria = cabecalho.index("categoria")
    idx_parcelado_csv = cabecalho.index("parcelado_csv")
    idx_parcelado_xlsx = cabecalho.index("parcelado_xlsx")

    conflito = next(linha for linha in dados if linha[idx_categoria] == "conflito_parcelamento")
    assert conflito[idx_parcelado_csv] == "nao"
    assert conflito[idx_parcelado_xlsx] == "sim"

    todos_os_valores = "".join(str(v) for linha in linhas for v in linha)
    assert "66666666666" not in todos_os_valores
    assert "77777777777" not in todos_os_valores


# --- dataset=matching-issues ------------------------------------------------


def test_exportar_matching_issues_csv_categorias(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/matching-issues",
        params={"format": "csv", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_csv(resposta.content)
    cabecalho, *dados = linhas
    idx_categoria = cabecalho.index("categoria")
    categorias = {linha[idx_categoria] for linha in dados}
    assert categorias == {"somente_csv", "somente_xlsx", "nome_ambiguo"}


def test_exportar_matching_issues_xlsx_categorias_e_nomes_ambiguos(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/matching-issues",
        params={"format": "xlsx", "csv_snapshot_id": csv_snap["id"], "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_xlsx(resposta.content)
    cabecalho, *dados = linhas
    idx_categoria = cabecalho.index("categoria")
    idx_ambiguo = cabecalho.index("nome_normalizado_ambiguo")

    categorias = {linha[idx_categoria] for linha in dados}
    assert categorias == {"somente_csv", "somente_xlsx", "nome_ambiguo"}

    linhas_ambiguas = [linha for linha in dados if linha[idx_categoria] == "nome_ambiguo"]
    assert len(linhas_ambiguas) == 2
    assert all(linha[idx_ambiguo] == "CARLOS SOUZA" for linha in linhas_ambiguas)


# --- dataset=entities ------------------------------------------------------


def test_exportar_entities_csv_respeita_filtro_query(cliente):
    _preparar_entidades(cliente)

    resposta = cliente.get("/api/exports/entities", params={"format": "csv", "query": "devedora"})
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_csv(resposta.content)
    cabecalho, *dados = linhas
    assert len(dados) == 1
    idx_nome = cabecalho.index("nome")
    assert dados[0][idx_nome] == "EMPRESA DEVEDORA LTDA"


def test_exportar_entities_csv_sem_filtro_ignora_paginacao(cliente):
    _preparar_entidades(cliente)

    resposta = cliente.get("/api/exports/entities", params={"format": "csv"})
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_csv(resposta.content)
    _cabecalho, *dados = linhas
    assert len(dados) == 3


def test_exportar_entities_xlsx_mascara_documento(cliente):
    _preparar_entidades(cliente)

    resposta = cliente.get("/api/exports/entities", params={"format": "xlsx", "query": "ana"})
    assert resposta.status_code == 200, resposta.text

    linhas = _linhas_xlsx(resposta.content)
    cabecalho, *dados = linhas
    assert len(dados) == 1
    idx_doc = cabecalho.index("documento_mascarado")
    assert dados[0][idx_doc] == "***.***.***-11"
    todos_os_valores = "".join(str(v) for linha in linhas for v in linha)
    assert "11111111111" not in todos_os_valores


# --- filtros de snapshot compartilhados com os endpoints JSON --------------


def test_exportar_usa_snapshots_mais_recentes_quando_ids_omitidos(cliente):
    _upload_csv(cliente, "csv_antigo.csv", ["ANTIGO;ANUIDADE;2025;0;01/01/2025;Débito"])
    _upload_xlsx(
        cliente,
        "xlsx_antigo_20260101_100000.xlsx",
        [_linha_xlsx_comparacao("11111111111", "ANTIGO XLSX", debitos=[{"ano": 2025, "tipo": "ANUIDADE"}])],
    )
    _upload_csv(cliente, "csv_recente.csv", ["RECENTE;ANUIDADE;2026;0;01/01/2026;Débito"])
    _upload_xlsx(
        cliente,
        "xlsx_recente_20260915_100000.xlsx",
        [_linha_xlsx_comparacao("22222222222", "RECENTE XLSX", debitos=[{"ano": 2026, "tipo": "ANUIDADE"}])],
    )

    resposta = cliente.get("/api/exports/matching-issues", params={"format": "csv"})
    assert resposta.status_code == 200, resposta.text
    assert "RECENTE" in resposta.text
    assert "ANTIGO" not in resposta.text


def test_exportar_snapshot_de_tipo_errado_retorna_422(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/comparisons",
        params={"csv_snapshot_id": xlsx_snap["id"], "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 422


def test_exportar_snapshot_inexistente_retorna_404(cliente):
    csv_snap, xlsx_snap = _preparar_par_completo(cliente)

    resposta = cliente.get(
        "/api/exports/comparisons",
        params={"csv_snapshot_id": 999999, "xlsx_snapshot_id": xlsx_snap["id"]},
    )
    assert resposta.status_code == 404


def test_exportar_dataset_invalido_retorna_422(cliente):
    resposta = cliente.get("/api/exports/nao-existe")
    assert resposta.status_code == 422


def test_exportar_formato_invalido_retorna_422(cliente):
    _preparar_entidades(cliente)

    resposta = cliente.get("/api/exports/entities", params={"format": "pdf"})
    assert resposta.status_code == 422
