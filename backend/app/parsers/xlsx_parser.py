from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

MAX_BLOCOS_DEBITO = 37
_EPOCH_EXCEL = date(1899, 12, 30)


@dataclass
class RegistroDebitoXlsx:
    ano_referencia: int
    tipo_debito: str | None
    data_vencimento: date | None
    valor_original: float | None
    valor_devido: float | None
    valor_total: float | None
    situacao_pagamento: str | None
    situacao_divida_ativa: str | None
    situacao_parcelamento: str | None


@dataclass
class RegistroXlsx:
    nome_original: str
    cpf_cnpj: str | None
    tipo_pessoa: str | None
    categoria: str | None
    subregiao: str | None
    situacao_registro: str | None
    debitos: list[RegistroDebitoXlsx] = field(default_factory=list)


def _texto(valor) -> str | None:
    if valor in (None, ""):
        return None
    return str(valor).strip()


def _numero(valor) -> float | None:
    if valor in (None, ""):
        return None
    return float(valor)


def _data(valor) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    # A célula normalmente já vem com formatação de data aplicada, e o openpyxl devolve
    # datetime nesse caso. Este branch é a rede de segurança para um número serial "cru"
    # (contagem de dias desde 1899-12-30, convenção do Excel), caso um export futuro do
    # CREFITO11 venha sem essa formatação.
    return _EPOCH_EXCEL + timedelta(days=int(float(valor)))


def parse_xlsx(conteudo: bytes) -> list[RegistroXlsx]:
    planilha = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    aba = planilha["Dados"]
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas)
    indice = {nome: i for i, nome in enumerate(cabecalho) if nome}

    registros = []
    for linha in linhas:
        if not linha or not linha[indice["NomeRazaoSocial"]]:
            continue

        registro = RegistroXlsx(
            nome_original=str(linha[indice["NomeRazaoSocial"]]).strip(),
            cpf_cnpj=_texto(linha[indice["CPFCNPJ"]]),
            tipo_pessoa=_texto(linha[indice["TipoPessoa"]]),
            categoria=_texto(linha[indice["Categoria"]]),
            subregiao=_texto(linha[indice["SubRegiao"]]),
            situacao_registro=_texto(linha[indice["SituacaoRegistro"]]),
        )

        for n in range(MAX_BLOCOS_DEBITO):
            campo_ano = f"Debitos.{n}.AnoReferencia"
            if campo_ano not in indice:
                break
            ano = linha[indice[campo_ano]]
            if not ano:
                continue
            registro.debitos.append(
                RegistroDebitoXlsx(
                    ano_referencia=int(ano),
                    tipo_debito=_texto(linha[indice[f"Debitos.{n}.DebitoTipo"]]),
                    data_vencimento=_data(linha[indice[f"Debitos.{n}.DataVencimento"]]),
                    valor_original=_numero(linha[indice[f"Debitos.{n}.ValorOriginal"]]),
                    valor_devido=_numero(linha[indice[f"Debitos.{n}.ValorDevido"]]),
                    valor_total=_numero(linha[indice[f"Debitos.{n}.ValorTotal"]]),
                    situacao_pagamento=_texto(linha[indice[f"Debitos.{n}.DebitoSituacaoPagamentoNome"]]),
                    situacao_divida_ativa=_texto(linha[indice[f"Debitos.{n}.DebitoSituacaoDividaAtivaNome"]]),
                    situacao_parcelamento=_texto(linha[indice[f"Debitos.{n}.DebitoSituacaoParcelamentoNome"]]),
                )
            )

        registros.append(registro)

    return registros
